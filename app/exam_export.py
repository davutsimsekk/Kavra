"""Local, deterministic printable exam exports; no model calls or remote resources."""
from html import escape
from io import BytesIO


def blocks(exam, answers=False):
    """One content model for Word and PDF; never include solutions in question papers."""
    yield "title", exam["name"]
    yield "subtitle", ("Cevap anahtarı" if answers else "Soru kâğıdı") + f" · {len(exam['questions'])} soru"
    if not answers:
        yield "text", "Ad Soyad: ........................................     Tarih: ...................."
    names = dict(zip(exam.get("sourceIds", []), exam.get("sourceNames", [])))
    for number, q in enumerate(exam["questions"], 1):
        yield "heading", f"Soru {number} · " + ("Çoktan seçmeli" if q["type"] == "mcq" else "Klasik")
        yield "prompt", q["prompt"]
        for i, option in enumerate(q.get("options", [])):
            yield "option", f"{chr(65+i)}) {option}"
        if answers:
            prefix = f"Doğru cevap: {chr(65+q['correctOption'])}) " if q["type"] == "mcq" else "Örnek cevap: "
            yield "answer", prefix + q["answer"]
            yield "text", q["explanation"]
            if q.get("rubric"):
                yield "label", "Değerlendirme ölçütleri"
                for item in q["rubric"]:
                    yield "text", "• " + item
            for item in q.get("evidence", []):
                yield "evidence", "Dayanak (" + names.get(item["sourceId"], "Kaynak") + "): " + item["quote"]
        elif q["type"] == "classic":
            yield "label", "Cevabınız"
            for _ in range(4):
                yield "blank", ". " * 65


def export_docx(exam, answers=False):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Subtitle", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    document.styles["Title"].font.size = Pt(22)
    document.styles["Heading 2"].font.size = Pt(12)
    document.core_properties.title = exam["name"]
    document.core_properties.author = "Kavra"
    document.core_properties.comments = ""
    for kind, text in blocks(exam, answers):
        style = {"title":"Title", "subtitle":"Subtitle", "heading":"Heading 2"}.get(kind, "Normal")
        paragraph = document.add_paragraph(text, style)
        paragraph.paragraph_format.widow_control = True
        if kind in ("title", "subtitle", "heading", "label"):
            paragraph.paragraph_format.keep_with_next = True
        if kind in ("answer", "label"):
            for run in paragraph.runs:
                run.bold = True
        if kind == "option":
            paragraph.paragraph_format.left_indent = Cm(.4)
        if kind == "blank":
            paragraph.paragraph_format.space_after = Pt(12)
        if kind == "evidence":
            for run in paragraph.runs:
                run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Sayfa ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_pdf(exam, answers=False):
    import pymupdf
    # Escape every user/model string: HTML is layout only, never executable content.
    groups = []
    parts = []
    for kind, text in blocks(exam, answers):
        if kind == "heading" and parts:
            groups.append("".join(parts))
            parts = []
        tag = "h1" if kind == "title" else "h2" if kind == "heading" else "p"
        parts.append(f'<{tag} class="{kind}">' + escape(text).replace("\n", "<br/>") + f"</{tag}>")
    if parts:
        groups.append("".join(parts))
    css = """
        body { font-family: sans-serif; font-size: 11pt; line-height: 1.35; }
        h1 { font-size: 22pt; margin: 0 0 8pt; }
        h2 { font-size: 12pt; margin: 18pt 0 7pt; page-break-after: avoid; }
        p { margin: 0 0 7pt; overflow-wrap: anywhere; }
        .subtitle { color: #555; margin-bottom: 18pt; }
        .option { margin-left: 12pt; }
        .answer, .label { font-weight: bold; }
        .label { page-break-after: avoid; }
        .blank { color: #aaa; margin-bottom: 12pt; }
        .evidence { font-size: 9pt; color: #444; }
    """
    page = pymupdf.paper_rect("a4")
    content = pymupdf.Rect(57, 51, page.width-57, page.height-55)
    output = BytesIO()
    writer = pymupdf.DocumentWriter(output)
    device = writer.begin_page(page)
    cursor = content.y0
    try:
        for html in groups:
            story = pymupdf.Story(html=html, user_css=css)
            available = pymupdf.Rect(content.x0, cursor, content.x1, content.y1)
            more, filled = story.place(available)
            if more and cursor > content.y0:
                # Move a question to the next page if it does not fit the remainder.
                # Oversized questions may still flow across multiple full pages.
                story.reset()
                writer.end_page()
                device = writer.begin_page(page)
                more, filled = story.place(content)
            while True:
                story.draw(device)
                if not more:
                    cursor = filled[3]
                    break
                writer.end_page()
                device = writer.begin_page(page)
                more, filled = story.place(content)
        writer.end_page()
    finally:
        writer.close()
    with pymupdf.open(stream=output.getvalue(), filetype="pdf") as document:
        document.set_metadata({"title":exam["name"], "author":"Kavra"})
        for i, sheet in enumerate(document, 1):
            sheet.insert_textbox(pymupdf.Rect(57, page.height-38, page.width-57, page.height-20),
                                 f"{i} / {len(document)}", fontsize=9, align=2)
        return document.tobytes(garbage=4, deflate=True)
