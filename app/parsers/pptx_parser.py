from pathlib import Path

from pptx import Presentation

from app.models import RawSection


def parse(path: Path) -> list[RawSection]:
    prs = Presentation(str(path))
    sections = []
    for i, slide in enumerate(prs.slides, start=1):
        title = None
        body_lines = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            is_title = shape == slide.shapes.title
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs).strip()
                if not line:
                    continue
                if is_title and title is None:
                    title = line
                else:
                    body_lines.append(line)

        notes = ""
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()

        text = "\n".join(body_lines)
        if notes:
            text += f"\n\n[Konuşmacı notu]\n{notes}"

        sections.append(RawSection(
            breadcrumb="",
            title=title or f"Slayt {i}",
            text=text.strip(),
            code_blocks=[],
            level=3,
        ))
    return [s for s in sections if s.text or s.title]
